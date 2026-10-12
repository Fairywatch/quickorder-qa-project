# -*- coding: utf-8 -*-
"""Genere docs/test-cases.xlsx a partir des resultats reels de la campagne."""
import io
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DOCS = os.path.join(ROOT, "docs")

results = json.load(io.open(os.path.join(DOCS, "evidence", "results.json"),
                            encoding="utf-8"))

# Priorite et rattachement au parcours critique
PRIO = {"Nominal": "P0", "Limite": "P2", "Erreur": "P1", "Regle metier": "P0"}
PARCOURS = {
    "N01": "PC6", "N02": "PC5", "N03": "PC1", "N04": "PC1", "N05": "PC6",
    "N06": "PC1", "N07": "PC1", "N08": "Supervision",
    "L01": "PC1", "L02": "PC1", "L03": "PC1", "L04": "PC1", "L05": "PC4",
    "L06": "PC4", "L07": "PC4", "L08": "PC2", "L09": "PC5", "L10": "PC5",
    "E01": "PC5", "E02": "PC5", "E03": "PC6", "E04": "PC6", "E05": "PC5",
    "E06": "PC4", "E07": "PC1", "E08": "PC1", "E09": "PC5", "E10": "PC5",
    "E11": "PC1", "E12": "PC1",
    "RM01": "PC1", "RM02": "PC3", "RM03": "PC2", "RM04": "PC2", "RM05": "PC1",
    "RM06": "PC1", "RM07": "PC4", "RM08": "PC6", "RM09": "PC6", "RM10": "PC4",
    "RM11": "PC6", "RM12": "PC7", "RM13": "PC6", "RM14": "PC2", "RM15": "PC1",
    "RM16": "PC1",
}
LIBELLE = {"Nominal": "Nominal", "Limite": "Limite", "Erreur": "Erreur",
           "Regle metier": "Regle metier"}

\
# Rendu sobre : pas de couleur ni de remplissage. L'en-tete se distingue par
# le gras, le statut par son texte. La grille reste, elle porte la lecture.
NO_FILL = PatternFill(fill_type=None)
HEAD_FILL = NO_FILL
HEAD_FONT = Font(bold=True, color="000000", size=11)
PASS_FILL = NO_FILL
PASS_FONT = Font(color="000000")
FAIL_FILL = NO_FILL
FAIL_FONT = Font(color="000000", bold=True)
ALT_FILL = NO_FILL
THIN = Side(style="thin", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLS = [
    ("ID", 9), ("Categorie", 14), ("Parcours", 11), ("Priorite", 9),
    ("Titre", 52), ("Objectif", 44), ("Preconditions", 42), ("Donnees de test", 40),
    ("Etapes", 52), ("Resultat attendu", 46), ("Resultat obtenu", 52),
    ("Statut", 10), ("Anomalie", 11),
]

wb = Workbook()

# ---------------------------------------------------------------- Cas de test
ws = wb.active
ws.title = "Cas de test"

for i, (name, width) in enumerate(COLS, start=1):
    c = ws.cell(row=1, column=i, value=name)
    c.fill = HEAD_FILL
    c.font = HEAD_FONT
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER
    ws.column_dimensions[get_column_letter(i)].width = width
ws.row_dimensions[1].height = 30

for idx, t in enumerate(results):
    r = idx + 2
    vals = [t["id"], LIBELLE[t["categorie"]], PARCOURS.get(t["id"], ""),
            PRIO[t["categorie"]], t["titre"], t["objectif"], t["preconditions"],
            t["donnees"], t["etapes"], t["attendu"], t["obtenu"], t["statut"],
            t["bug"] or "-"]
    for i, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = BORDER
        if idx % 2 == 1:
            c.fill = ALT_FILL
    sc = ws.cell(row=r, column=12)
    if t["statut"] == "Pass":
        sc.fill, sc.font = PASS_FILL, PASS_FONT
    elif t["statut"] == "Fail":
        sc.fill, sc.font = FAIL_FILL, FAIL_FONT
    sc.alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=r, column=1).font = Font(bold=True)
    ws.row_dimensions[r].height = 58

ws.freeze_panes = "E2"
ws.auto_filter.ref = "A1:M%d" % (len(results) + 1)

# ---------------------------------------------------------------- Synthese
ws2 = wb.create_sheet("Synthese")
ws2.column_dimensions["A"].width = 34
ws2.column_dimensions["B"].width = 14
ws2.column_dimensions["C"].width = 14
ws2.column_dimensions["D"].width = 14
ws2.column_dimensions["E"].width = 14

title = ws2.cell(row=1, column=1, value="Synthese de campagne — QuickOrder V3.1")
title.font = Font(bold=True, size=14, color="000000")
ws2.cell(row=2, column=1, value="Execution du 22/09/2026 — base reinitialisee avant campagne")

cats = ["Nominal", "Limite", "Erreur", "Regle metier"]
r0 = 4
for i, h in enumerate(["Categorie", "Total", "Pass", "Fail", "Taux de reussite"], start=1):
    c = ws2.cell(row=r0, column=i, value=h)
    c.fill, c.font = HEAD_FILL, HEAD_FONT
    c.alignment = Alignment(horizontal="center")
    c.border = BORDER

for j, cat in enumerate(cats):
    rr = r0 + 1 + j
    sub = [t for t in results if t["categorie"] == cat]
    p = sum(1 for t in sub if t["statut"] == "Pass")
    f = sum(1 for t in sub if t["statut"] == "Fail")
    taux = (float(p) / len(sub) * 100) if sub else 0
    for i, v in enumerate([LIBELLE[cat], len(sub), p, f, "%.0f %%" % taux], start=1):
        c = ws2.cell(row=rr, column=i, value=v)
        c.border = BORDER
        if i > 1:
            c.alignment = Alignment(horizontal="center")
    if f:
        ws2.cell(row=rr, column=4).fill = FAIL_FILL
        ws2.cell(row=rr, column=4).font = FAIL_FONT

rt = r0 + 1 + len(cats)
tp = sum(1 for t in results if t["statut"] == "Pass")
tf = sum(1 for t in results if t["statut"] == "Fail")
for i, v in enumerate(["TOTAL", len(results), tp, tf,
                       "%.0f %%" % (float(tp) / len(results) * 100)], start=1):
    c = ws2.cell(row=rt, column=i, value=v)
    c.font = Font(bold=True)
    c.border = BORDER
    c.fill = NO_FILL
    if i > 1:
        c.alignment = Alignment(horizontal="center")

ws2.cell(row=rt + 2, column=1, value="Statuts complementaires").font = Font(bold=True)
ws2.cell(row=rt + 3, column=1, value="Blocked")
ws2.cell(row=rt + 3, column=2,
         value=sum(1 for t in results if t["statut"] == "Blocked")).alignment = \
    Alignment(horizontal="center")
ws2.cell(row=rt + 4, column=1, value="Skipped")
ws2.cell(row=rt + 4, column=2,
         value=sum(1 for t in results if t["statut"] == "Skipped")).alignment = \
    Alignment(horizontal="center")

concl = ws2.cell(row=rt + 6, column=1,
                 value="Recommandation : NO-GO — 13 anomalies dont 1 bloquante et "
                       "4 critiques sur les parcours de paiement et de stock.")
concl.font = Font(bold=True, color="000000", size=12)

# ---------------------------------------------------------------- Anomalies
ws3 = wb.create_sheet("Anomalies")
bugs = {}
for t in results:
    if t["statut"] == "Fail" and t["bug"]:
        bugs.setdefault(t["bug"], []).append(t["id"])

SEV = {
    "BUG-005": ("Bloquante", "P0", "Paiement PayPal : deux commandes creees pour un seul paiement"),
    "BUG-002": ("Critique", "P1", "Un restaurant ferme accepte les commandes"),
    "BUG-003": ("Critique", "P1", "Un produit en rupture de stock reste commandable"),
    "BUG-006": ("Critique", "P1", "Paiement en especes insuffisant accepte et marque paye"),
    "BUG-010": ("Critique", "P1", "Le sous-total transmis par le client n'est pas verifie"),
    "BUG-001": ("Majeure", "P1", "Montant minimum de commande jamais applique"),
    "BUG-004": ("Majeure", "P1", "Le stock devient negatif"),
    "BUG-007": ("Majeure", "P2", "Coupon ATLAS15 annonce a 15 %, applique a 10 %"),
    "BUG-011": ("Majeure", "P2", "Avis possible sans avoir commande"),
    "BUG-012": ("Majeure", "P2", "Transitions de statut impossibles acceptees"),
    "BUG-008": ("Majeure", "P2", "Filtre frais de livraison inverse"),
    "BUG-009": ("Mineure", "P3", "La recherche ignore la description"),
    "BUG-013": ("Mineure", "P3", "L'API expose les restaurants fermes"),
}
for i, (h, w) in enumerate([("ID", 12), ("Titre", 62), ("Severite", 13),
                            ("Priorite", 10), ("Cas de test", 18)], start=1):
    c = ws3.cell(row=1, column=i, value=h)
    c.fill, c.font = HEAD_FILL, HEAD_FONT
    c.alignment = Alignment(horizontal="center")
    c.border = BORDER
    ws3.column_dimensions[get_column_letter(i)].width = w

ordre = {"Bloquante": 0, "Critique": 1, "Majeure": 2, "Mineure": 3}
for j, bid in enumerate(sorted(SEV, key=lambda k: (ordre[SEV[k][0]], k))):
    sev, prio, titre = SEV[bid]
    rr = j + 2
    for i, v in enumerate([bid, titre, sev, prio, ", ".join(sorted(bugs.get(bid, [])))],
                          start=1):
        c = ws3.cell(row=rr, column=i, value=v)
        c.border = BORDER
        c.alignment = Alignment(vertical="center", wrap_text=(i == 2))
    sc = ws3.cell(row=rr, column=3)
    sc.alignment = Alignment(horizontal="center")
    if sev in ("Bloquante", "Critique"):
        sc.fill, sc.font = FAIL_FILL, FAIL_FONT
    elif sev == "Majeure":
        sc.fill = NO_FILL
ws3.freeze_panes = "A2"

out = os.path.join(DOCS, "test-cases.xlsx")
wb.save(out)
print("ecrit : %s (%d cas, 3 feuilles)" % (out, len(results)))
